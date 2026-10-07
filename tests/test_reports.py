"""Тесты выгрузки отчётов в Excel (день 9)."""

import os
import tempfile
import unittest

from openpyxl import load_workbook

from services import apartments, auth, finance, reports
from tests.helpers import BaseTest, apartment_data, chairman_data


class ReportsTest(BaseTest):
    def setUp(self):
        super().setUp()
        auth.register_chairman(chairman_data())
        apartments.save_apartment(apartment_data("1", "40"))
        apartments.save_apartment(apartment_data("2", "30"))
        finance.charge_month("2026-09")
        apartment_id = apartments.get_apartments()[1]["apartment_id"]
        finance.add_payment(apartment_id, "975", "05.09.2026", "Перевод")
        handle, self.file_path = tempfile.mkstemp(suffix=".xlsx")
        os.close(handle)

    def tearDown(self):
        os.remove(self.file_path)
        super().tearDown()

    def test_debtors_report_contains_only_debtors(self):
        count = reports.export_debtors(self.file_path)
        self.assertEqual(count, 1)
        sheet = load_workbook(self.file_path).active
        self.assertEqual(sheet.cell(row=2, column=1).value, 1)
        self.assertEqual(sheet.cell(row=2, column=4).value, 1300.0)


if __name__ == "__main__":
    unittest.main()
